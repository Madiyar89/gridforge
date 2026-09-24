"""Накопительный Excel-реестр оценки уязвимости — перенос из NetOpsHub
(backend/app/modules/reports/vuln_register.py), тот же формат бланка и та
же идея "реестр дописывается, не пересоздаётся", но источник данных —
не отдельные .xml-файлы на диске, а уже сохранённые VulnScan/VulnScanHost
в БД (см. app/vuln_scan_engine.py). Один файл на группу:
DATA_DIR/vuln-register/{group_id}/register.xlsx.

Идемпотентность: id уже внесённых VulnScan хранится на скрытом листе
того же .xlsx (_ingested_scans) — повторный вызов ingest_scan на тот же
scan.id не плодит дубли строк (актуально, если фоновая задача случайно
перевыполнится).

Каждый ЗАВЕРШЁННЫЙ скан оставляет минимум одну строку — либо находки,
либо одна явная строка "не выявлено": так по файлу видно "проверяли,
чисто", а не тишину, которую нельзя отличить от "не проверяли вовсе"
(тот же принцип, что в NetOpsHub после правки от 2026-08-25).

Формат — официальный бланк "Отчёт о проведении оценки уязвимости сетевых
ресурсов ... Республики Казахстан" (10-11 колонок, см. HEADERS ниже)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import desc
from sqlalchemy.orm import Session

from app.db import DATA_DIR
from app.models import VulnScan, VulnScanStatus, VULN_SCAN_PROFILE_LABELS

REGISTER_DIR = DATA_DIR / "vuln-register"

MAIN_SHEET = "Оценка уязвимости"
INGESTED_SHEET = "_ingested_scans"

HEADERS = [
    "Дата проведения оценки уязвимости ресурсов",
    "Ф.И.О. и должность работника",
    "Тип проверки",
    "Описание уязвимости",
    "Уровень критичности",
    "Анализ уязвимости",
    "Перечень ресурсов, подтвержденных уязвимостей",
    "Устранение уязвимости",
    "Результат устранения уязвимости",
    "Проверка уязвимости",
    "Подпись работника управления информационной безопасности",
]
AUTO_COLUMNS = {0, 2, 3, 4, 5, 6}
SEVERITY_LABEL = {"vulnerable": "Критично", "likely": "Вероятно", "unknown": "Требует проверки"}
SEVERITY_RECOMMENDATION = {
    "vulnerable": "Устранить согласно рекомендациям производителя ПО/прошивки; при невозможности сразу — ограничить сетевой доступ к сервису.",
    "likely": "Проверить вручную и устранить при подтверждении; при невозможности сразу — ограничить сетевой доступ к сервису.",
    "unknown": "Автоматическая проверка не смогла подтвердить или опровергнуть — требуется ручная проверка.",
}


def register_path(group_id: int) -> Path:
    return REGISTER_DIR / str(group_id) / "register.xlsx"


def _new_workbook(group_name: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = MAIN_SHEET

    title_font = Font(name="Times New Roman", size=13, bold=True)
    header_font = Font(name="Times New Roman", size=10, bold=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title_cell = ws.cell(
        row=1,
        column=1,
        value=f"Отчёт о проведении оценки уязвимости сетевых ресурсов {group_name} Республики Казахстан",
    )
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    for col, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=text)
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 46

    for col, width in enumerate([14, 24, 26, 36, 14, 26, 32, 20, 20, 18, 28], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ingested_ws = wb.create_sheet(INGESTED_SHEET)
    ingested_ws.append(["scan_id"])
    ingested_ws.sheet_state = "hidden"
    return wb


def _write_row(ws, row_idx: int, row_values: list, body_font: Font, auto_fill: PatternFill, wrap: Alignment) -> None:
    for col, value in enumerate(row_values, start=1):
        cell = ws.cell(row=row_idx, column=col, value=value)
        cell.font = body_font
        cell.alignment = wrap
        if (col - 1) in AUTO_COLUMNS:
            cell.fill = auto_fill


def ingest_scan(db: Session, scan: VulnScan) -> bool:
    """Дописывает результат VulnScan в реестр его группы, если он ещё не
    вносился — находки построчно, а если находок нет ни на одном хосте —
    одна явная строка "не выявлено". Возвращает True, если что-то реально
    добавлено (используется только для тестов/логов — main.py на это не
    полагается)."""
    path = register_path(scan.group_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(path) if path.exists() else _new_workbook(scan.group.name)

    ingested_ws = wb[INGESTED_SHEET]
    already = {row[0] for row in ingested_ws.iter_rows(min_row=2, values_only=True) if row[0] is not None}
    if scan.id in already:
        wb.close()
        return False

    ws = wb[MAIN_SHEET]
    profile_label = VULN_SCAN_PROFILE_LABELS.get(scan.profile.value, scan.profile.value)
    scan_date = scan.started_at.strftime("%Y-%m-%d %H:%M") if scan.started_at else ""
    body_font = Font(name="Times New Roman", size=10)
    auto_fill = PatternFill("solid", fgColor="E3F2E8")
    wrap = Alignment(wrap_text=True, vertical="top")

    hosts_up = [h for h in scan.hosts if h.state == "up"]
    row_idx = ws.max_row + 1
    found_any = False
    for host in hosts_up:
        resource = host.address
        if host.hostname:
            resource += f" ({host.hostname})"
        for finding in host.findings or []:
            severity = finding.get("severity")
            if severity == "info":
                continue
            found_any = True
            description = finding.get("summary", "")
            cves = finding.get("cves") or []
            if cves:
                description = f"{description} ({', '.join(cves)})" if description else ", ".join(cves)
            analysis = SEVERITY_RECOMMENDATION.get(severity, "")

            _write_row(
                ws,
                row_idx,
                [
                    scan_date,
                    scan.responsible or "",
                    profile_label,
                    description,
                    SEVERITY_LABEL.get(severity, severity or ""),
                    analysis,
                    f"{resource} — {finding.get('port', '')}",
                    "",
                    "",
                    "",
                    "",
                ],
                body_font,
                auto_fill,
                wrap,
            )
            row_idx += 1

    if not found_any:
        _write_row(
            ws,
            row_idx,
            [
                scan_date,
                scan.responsible or "",
                profile_label,
                "не выявлено",
                "низкий",
                f"не обнаружено — хостов найдено {len(hosts_up)}",
                "не обнаружен",
                "нет необходимости",
                "нет необходимости",
                "не обнаружено",
                "",
            ],
            body_font,
            auto_fill,
            wrap,
        )
        row_idx += 1

    ingested_ws.append([scan.id])
    wb.save(path)
    wb.close()
    return True


def _existing_profile_keys(db: Session, group_id: int) -> set[str]:
    return {
        row[0]
        for row in db.query(VulnScan.profile)
        .filter(VulnScan.group_id == group_id, VulnScan.status == VulnScanStatus.done)
        .distinct()
        .all()
    }


def missing_profile_rows(db: Session, group_id: int) -> list[list]:
    """Строки-заглушки для профилей без единого завершённого скана у этой
    группы — "выполняется" или "не запускался". НЕ пишутся в сам .xlsx —
    вычисляются заново на каждый показ/скачивание."""
    existing = {p.value for p in _existing_profile_keys(db, group_id)}
    running_profiles = {
        row[0].value
        for row in db.query(VulnScan.profile)
        .filter(VulnScan.group_id == group_id, VulnScan.status == VulnScanStatus.running)
        .distinct()
        .all()
    }

    rows = []
    for key, label in VULN_SCAN_PROFILE_LABELS.items():
        if key in existing:
            continue
        if key in running_profiles:
            rows.append(["", "", label, "скан выполняется", "—", "результат появится после завершения", "—", "—", "—", "—", ""])
        else:
            rows.append(["", "", label, "скан не запускался", "—", "профиль ещё ни разу не использовался для этой группы", "—", "—", "—", "—", ""])
    return rows


def read_persisted_rows(group_id: int) -> list[tuple]:
    """Реальные (не заглушки) строки реестра — для показа на сайте."""
    path = register_path(group_id)
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb[MAIN_SHEET]
        return [row for row in ws.iter_rows(min_row=3, values_only=True) if any(row)]
    finally:
        wb.close()


def build_excel_with_missing(db: Session, group_id: int, group_name: str) -> Workbook:
    """Реестр (существующий или новый пустой) + строки-заглушки поверх —
    для скачивания. Заглушки НЕ сохраняются обратно в register.xlsx на
    диске — только в возвращаемом объекте."""
    path = register_path(group_id)
    wb = load_workbook(path) if path.exists() else _new_workbook(group_name)
    ws = wb[MAIN_SHEET]
    missing = missing_profile_rows(db, group_id)
    if missing:
        note_font = Font(name="Times New Roman", size=10, italic=True, color="7A7A72")
        note_fill = PatternFill("solid", fgColor="F2F0E8")
        wrap = Alignment(wrap_text=True, vertical="top")
        row_idx = ws.max_row + 1
        for row_values in missing:
            _write_row(ws, row_idx, row_values, note_font, note_fill, wrap)
            row_idx += 1
    return wb


def last_scan_summary(db: Session, group_id: int) -> list[dict]:
    """Последний скан по каждому профилю (для карточек на странице) —
    статус, когда, сколько находок."""
    result = []
    for key, label in VULN_SCAN_PROFILE_LABELS.items():
        scan = (
            db.query(VulnScan)
            .filter(VulnScan.group_id == group_id, VulnScan.profile == key)
            .order_by(desc(VulnScan.started_at))
            .first()
        )
        if scan is None:
            result.append({"profile": key, "label": label, "scan": None})
            continue
        findings_count = sum(
            1 for h in scan.hosts for f in (h.findings or []) if f.get("severity") != "info"
        )
        result.append(
            {
                "profile": key,
                "label": label,
                "scan": {
                    "id": scan.id,
                    "status": scan.status.value,
                    "started_at": scan.started_at,
                    "finished_at": scan.finished_at,
                    "error": scan.error,
                    "host_count": len(scan.hosts),
                    "findings_count": findings_count,
                },
            }
        )
    return result
